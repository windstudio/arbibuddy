from pathlib import Path
import re
from typing import Any, Mapping

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _set_run_font(run, *, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def _set_style_font(style, *, name: str, size: float) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def _page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_field(instruction_text: str, display_text: str) -> None:
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction_text)
        display_run = paragraph.add_run(display_text)
        _set_run_font(display_run, name="Times New Roman", size=9)
        paragraph._p.remove(display_run._r)
        field.append(display_run._r)
        paragraph._p.append(field)

    add_field("PAGE", "1")
    separator = paragraph.add_run(" / ")
    _set_run_font(separator, name="Times New Roman", size=9)
    add_field("NUMPAGES", "1")


def _set_style_list_indentation(style, *, first_line_chars: int) -> None:
    properties = style.element.get_or_add_pPr()
    indentation = properties.find(qn("w:ind"))
    if indentation is None:
        indentation = OxmlElement("w:ind")
        properties.append(indentation)
    for attribute in (
        "firstLine", "firstLineChars", "left", "leftChars", "hanging", "hangingChars",
    ):
        indentation.attrib.pop(qn(f"w:{attribute}"), None)
    indentation.set(qn("w:firstLineChars"), str(first_line_chars))


def _set_numbering_list_indentation(
    document,
    style,
    *,
    first_line_chars: int,
) -> None:
    numbering = document.part.numbering_part.element
    num_id = style.element.pPr.numPr.numId.val
    source_num = numbering.find(
        qn("w:num") + f"[@{qn('w:numId')}='{num_id}']"
    )
    if source_num is None:
        raise ValueError("Word列表编号定义缺失")
    abstract_id_node = source_num.find(qn("w:abstractNumId"))
    if abstract_id_node is None:
        raise ValueError("Word列表抽象编号引用缺失")
    abstract_id = abstract_id_node.get(qn("w:val"))
    abstract = numbering.find(
        qn("w:abstractNum") + f"[@{qn('w:abstractNumId')}='{abstract_id}']"
    )
    if abstract is None:
        raise ValueError("Word列表抽象编号定义缺失")
    for level in abstract.findall(qn("w:lvl")):
        properties = level.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            level.append(properties)
        indentation = properties.find(qn("w:ind"))
        if indentation is None:
            indentation = OxmlElement("w:ind")
            properties.append(indentation)
        for attribute in (
            "firstLine", "firstLineChars", "left", "leftChars",
            "hanging", "hangingChars",
        ):
            indentation.attrib.pop(qn(f"w:{attribute}"), None)
        indentation.set(qn("w:firstLineChars"), str(first_line_chars))


def _new_legal_document(*, landscape: bool = False):
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(29.7) if landscape else Cm(21)
    section.page_height = Cm(21) if landscape else Cm(29.7)
    section.orientation = 1 if landscape else 0
    section.top_margin = Cm(2.0) if landscape else Cm(2.6)
    section.bottom_margin = Cm(2.0) if landscape else Cm(2.4)
    section.left_margin = Cm(1.8) if landscape else Cm(2.8)
    section.right_margin = Cm(1.8) if landscape else Cm(2.6)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = document.styles["Normal"]
    _set_style_font(normal, name="仿宋", size=12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.widow_control = True

    body_style = document.styles.add_style(
        "ArbiBuddy Legal Body", WD_STYLE_TYPE.PARAGRAPH
    )
    _set_style_font(body_style, name="仿宋", size=12)
    body_style.paragraph_format.space_before = Pt(0)
    body_style.paragraph_format.space_after = Pt(0)
    body_style.paragraph_format.line_spacing = 1.5
    body_style.paragraph_format.widow_control = True
    for style_name in ("List Number", "List Bullet"):
        list_style = document.styles[style_name]
        _set_style_list_indentation(list_style, first_line_chars=200)
        _set_numbering_list_indentation(
            document,
            list_style,
            first_line_chars=200,
        )
        list_style.paragraph_format.line_spacing = 1.5
        list_style.paragraph_format.space_before = Pt(0)
        list_style.paragraph_format.space_after = Pt(0)

    title_style = document.styles.add_style("ArbiBuddy Legal Title", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(title_style, name="黑体", size=22)
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(0)
    title_style.paragraph_format.keep_with_next = True
    title_style.paragraph_format.widow_control = True

    signature_style = document.styles.add_style("ArbiBuddy Legal Signature", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(signature_style, name="仿宋", size=12)
    signature_style.paragraph_format.space_after = Pt(0)
    signature_style.paragraph_format.line_spacing = 1.5
    signature_style.paragraph_format.keep_together = True
    signature_style.paragraph_format.widow_control = True

    footer = section.footer.paragraphs[0]
    _page_field(footer)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    return document


def _set_character_indentation(
    paragraph,
    *,
    first_line_chars: int | None = None,
    left_chars: int | None = None,
    hanging_chars: int | None = None,
) -> None:
    properties = paragraph._p.get_or_add_pPr()
    indentation = properties.find(qn("w:ind"))
    if indentation is None:
        indentation = OxmlElement("w:ind")
        properties.append(indentation)
    for attribute in (
        "firstLine", "firstLineChars", "left", "leftChars", "hanging", "hangingChars",
    ):
        indentation.attrib.pop(qn(f"w:{attribute}"), None)
    if first_line_chars is not None:
        indentation.set(qn("w:firstLineChars"), str(first_line_chars))
    if left_chars is not None:
        indentation.set(qn("w:leftChars"), str(left_chars))
    if hanging_chars is not None:
        indentation.set(qn("w:hangingChars"), str(hanging_chars))


def _add_blank_paragraph(document) -> None:
    document.add_paragraph()


def _add_title(document, text: str) -> None:
    title = document.add_paragraph(style="ArbiBuddy Legal Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    run = title.add_run(text)
    _set_run_font(run, name="黑体", size=22, bold=True)
    _add_blank_paragraph(document)


def _add_heading(document, text: str):
    paragraph = document.add_paragraph(style="ArbiBuddy Legal Body")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    _set_run_font(run, name="黑体", size=12, bold=False)
    return paragraph


def _add_body_paragraph(document, text: str, *, keep_with_next: bool = False):
    paragraph = document.add_paragraph(style="ArbiBuddy Legal Body")
    _set_character_indentation(paragraph, first_line_chars=200)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.widow_control = True
    paragraph.add_run(text)
    return paragraph


def _start_numbered_list(document) -> int:
    """Create a list instance that starts at one without changing other lists."""
    numbering = document.part.numbering_part.element
    style_num_id = document.styles['List Number'].element.pPr.numPr.numId.val
    source = numbering.find(qn('w:num') + f"[@{qn('w:numId')}='{style_num_id}']")
    abstract = source.find(qn('w:abstractNumId')) if source is not None else None
    if abstract is None:
        raise ValueError('Word列表抽象编号引用缺失')
    num_id = max(int(item.get(qn('w:numId'))) for item in numbering.findall(qn('w:num'))) + 1
    number = OxmlElement('w:num')
    number.set(qn('w:numId'), str(num_id))
    reference = OxmlElement('w:abstractNumId')
    reference.set(qn('w:val'), abstract.get(qn('w:val')))
    number.append(reference)
    override = OxmlElement('w:lvlOverride')
    override.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:startOverride')
    start.set(qn('w:val'), '1')
    override.append(start)
    number.append(override)
    numbering.append(number)
    return num_id


def _add_numbered_paragraph(document, text: str, *, num_id: int | None = None):
    paragraph = document.add_paragraph(style="List Number")
    if num_id is not None:
        properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_ilvl().val = 0
        properties.get_or_add_numId().val = num_id
    _set_character_indentation(paragraph, first_line_chars=200)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(text)
    return paragraph


def _add_closing(document, recipient: str) -> None:
    closing = document.add_paragraph(style="ArbiBuddy Legal Body")
    _set_character_indentation(closing, first_line_chars=200)
    closing.paragraph_format.keep_with_next = True
    closing.add_run("此致")
    recipient_paragraph = document.add_paragraph(style="ArbiBuddy Legal Body")
    recipient_paragraph.paragraph_format.keep_with_next = True
    recipient_paragraph.add_run(recipient)


def _iter_paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _save_legal_document(document, path: Path) -> None:
    """Materialize effective fonts so office-suite reserialization cannot lose them."""
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    seen: set[int] = set()
    for container in containers:
        for paragraph in _iter_paragraphs(container):
            paragraph_id = id(paragraph._p)
            if paragraph_id in seen:
                continue
            seen.add(paragraph_id)
            style_font = paragraph.style.font
            for run in paragraph.runs:
                name = run.font.name or style_font.name or "仿宋"
                size = (
                    run.font.size.pt
                    if run.font.size is not None
                    else style_font.size.pt
                    if style_font.size is not None
                    else 12
                )
                _set_run_font(
                    run,
                    name=name,
                    size=size,
                    bold=bool(
                        run.bold
                        if run.bold is not None
                        else style_font.bold
                    ),
                )
    document.save(path)


def _set_cell_width(cell, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: tuple[int, ...]) -> None:
    properties = table._tbl.tblPr
    for name, attributes in (
        ("w:tblW", {"w:w": str(sum(widths_dxa)), "w:type": "dxa"}),
        ("w:tblInd", {"w:w": "0", "w:type": "dxa"}),
        ("w:tblLayout", {"w:type": "fixed"}),
    ):
        element = properties.first_child_found_in(name)
        if element is None:
            element = OxmlElement(name)
            properties.append(element)
        for key, value in attributes.items():
            element.set(qn(key), value)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_dxa in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width_dxa))
        grid.append(column)
    for row in table.rows:
        for cell, width_dxa in zip(row.cells, widths_dxa, strict=True):
            _set_cell_width(cell, width_dxa)
            margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell._tc.get_or_add_tcPr().append(margins)
            for side, value in (("top", 80), ("start", 100), ("bottom", 80), ("end", 100)):
                margin = margins.find(qn(f"w:{side}"))
                if margin is None:
                    margin = OxmlElement(f"w:{side}")
                    margins.append(margin)
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


# All registered builders cross the same XML 1.0 field-binding seam before
# creating a document or its parent output directory.
