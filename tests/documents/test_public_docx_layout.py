from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from docx.oxml.ns import qn

from scripts.documents.public import _build_docx


def _indent_value(paragraph, name: str) -> str | None:
    properties = paragraph._p.pPr
    if properties is None:
        return None
    indentation = properties.find(qn("w:ind"))
    return None if indentation is None else indentation.get(qn(f"w:{name}"))


class PublicDocxLayoutTests(unittest.TestCase):
    def test_arbitration_application_uses_semantic_sections_and_real_numbering(self):
        request = {
            "document_type": "arbitration_application",
            "title": "劳动人事争议仲裁申请书",
            "locked_bindings": [
                {"kind": "document_heading", "rendered_value": "劳动人事争议仲裁申请书"},
                {"kind": "claim", "rendered_value": "拖欠工资"},
                {"kind": "claim", "rendered_value": "季度奖金"},
                {"kind": "amount", "rendered_value": "36000.00 元"},
                {"kind": "date", "rendered_value": "2025-08-31"},
                {"kind": "signature", "rendered_value": "申请人：{{申请人姓名}}"},
            ],
            "sections": [
                {
                    "heading": "申请人",
                    "full_text": "申请人：{{申请人姓名}}，联系电话：{{申请人联系电话}}。",
                },
                {
                    "heading": "被申请人",
                    "full_text": "被申请人：{{用人单位全称}}。",
                },
                {
                    "heading": "仲裁请求",
                    "full_text": (
                        "一、请求支付拖欠工资；\n"
                        "二、请求支付季度奖金；\n"
                        "三、请求支付加班工资。"
                    ),
                },
                {
                    "heading": "事实与理由",
                    "full_text": "第一段事实。\n\n第二段事实。\n\n第三段事实。",
                },
                {
                    "heading": "落款",
                    "full_text": (
                        "此致\n{{劳动人事争议仲裁委员会名称}}\n\n"
                        "申请人（签名）：{{申请人姓名}}\n\n"
                        "联系电话：{{申请人联系电话}}\n\n"
                        "日期：{{申请日期}}"
                    ),
                },
            ],
        }

        with TemporaryDirectory() as temporary_root:
            path = Path(temporary_root) / "application.docx"
            _build_docx(path, request)
            document = Document(path)

        paragraphs = document.paragraphs
        texts = [paragraph.text for paragraph in paragraphs]
        self.assertEqual(texts[:2], ["劳动人事争议仲裁申请书", ""])
        for leaked in ("拖欠工资", "季度奖金", "36000.00 元", "2025-08-31"):
            self.assertNotIn(leaked, texts)
        for structural_heading in ("申请人", "被申请人", "落款"):
            self.assertNotIn(structural_heading, texts)

        numbered = [paragraph for paragraph in paragraphs if paragraph.style.name == "List Number"]
        self.assertEqual(
            [paragraph.text for paragraph in numbered],
            ["请求支付拖欠工资；", "请求支付季度奖金；", "请求支付加班工资。"],
        )
        fact_paragraphs = [paragraph for paragraph in paragraphs if paragraph.text.endswith("段事实。")]
        self.assertEqual(len(fact_paragraphs), 3)
        self.assertEqual(
            [_indent_value(paragraph, "firstLineChars") for paragraph in fact_paragraphs],
            ["200", "200", "200"],
        )

        signature = [
            paragraph
            for paragraph in paragraphs
            if paragraph.style.name == "ArbiBuddy Legal Signature"
        ]
        self.assertEqual(len(signature), 3)
        self.assertTrue(all(_indent_value(paragraph, "leftChars") == "1600" for paragraph in signature))
        self.assertEqual(texts[texts.index(signature[0].text) - 1], "")

    def test_evidence_catalog_renders_only_title_and_five_column_table(self):
        request = {
            "document_type": "evidence_catalog",
            "title": "证据目录",
            "locked_bindings": [
                {"kind": "document_heading", "rendered_value": "证据目录"},
                {"kind": "evidence", "rendered_value": "工资流水｜银行流水｜证明工资差额"},
            ],
            "sections": [
                {
                    "heading": "标题",
                    "full_text": (
                        "证据目录\n\n提交人：{{申请人姓名}}\n"
                        "对方当事人：{{用人单位全称}}\n案由：工资争议"
                    ),
                },
                {
                    "heading": "五列表格",
                    "full_text": (
                        "| 编号 | 证据名称 | 证据内容 | 证明目的 | 页码 |\n"
                        "| --- | --- | --- | --- | --- |\n"
                        "| 1 | 工资流水 | 银行流水 | 证明工资差额 |  |"
                    ),
                },
            ],
        }

        with TemporaryDirectory() as temporary_root:
            path = Path(temporary_root) / "evidence.docx"
            _build_docx(path, request)
            document = Document(path)

        self.assertEqual([paragraph.text for paragraph in document.paragraphs], ["证据目录", ""])
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(
            [cell.text for cell in document.tables[0].rows[0].cells],
            ["编号", "证据名称", "证据内容", "证明目的", "页码"],
        )

    def test_forced_termination_notice_formats_sections_placeholders_and_signature(self):
        request = {
            "document_type": "forced_termination_notice",
            "title": "被迫解除劳动合同通知书",
            "locked_bindings": [
                {"kind": "document_heading", "rendered_value": "被迫解除劳动合同通知书"},
                {"kind": "party", "rendered_value": "{{用人单位全称}}"},
                {"kind": "party", "rendered_value": "{{劳动者姓名}}"},
                {"kind": "claim", "rendered_value": "未及时足额支付劳动报酬"},
                {"kind": "evidence", "rendered_value": "工资发放银行流水"},
                {"kind": "date", "rendered_value": "{{解除生效日期}}"},
                {"kind": "signature", "rendered_value": "{{劳动者签名}}"},
            ],
            "sections": [
                {
                    "heading": "收件人",
                    "full_text": (
                        "被迫解除劳动合同通知书\n\n致：{{用人单位全称}}\n\n"
                        "本人{{劳动者姓名}}，现向你单位发出本通知。"
                    ),
                },
                {
                    "heading": "解除意思表示",
                    "full_text": "本人正式通知你单位解除劳动合同。\n\n解除自{{解除生效日期}}生效。",
                },
                {
                    "heading": "解除理由",
                    "full_text": (
                        "一、你单位连续欠付工资。\n\n"
                        "二、工资差额至今未补足。\n\n"
                        "三、上述事实有银行流水证明。\n\n"
                        "据此，本人依法解除劳动合同。"
                    ),
                },
                {
                    "heading": "结算与手续",
                    "full_text": "请依法办理结算。\n\n一、结清工资。\n\n二、出具解除证明。",
                },
                {
                    "heading": "落款",
                    "full_text": (
                        "通知人（签名）：{{劳动者签名}}\n\n"
                        "日期：{{签署日期}}\n\n"
                        "送达方式：{{送达方式}}，并保留送达凭证。"
                    ),
                },
            ],
        }

        with TemporaryDirectory() as temporary_root:
            path = Path(temporary_root) / "notice.docx"
            _build_docx(path, request)
            document = Document(path)

        paragraphs = document.paragraphs
        texts = [paragraph.text for paragraph in paragraphs]
        for structural_heading in ("收件人", "解除意思表示", "落款"):
            self.assertNotIn(structural_heading, texts)
        self.assertEqual(texts.count("被迫解除劳动合同通知书"), 1)
        self.assertFalse(any("{{" in text or "}}" in text for text in texts))
        self.assertTrue(any("【待补：用人单位全称】" in text for text in texts))

        numbered = [paragraph for paragraph in paragraphs if paragraph.style.name == "List Number"]
        self.assertEqual(
            [paragraph.text for paragraph in numbered],
            ["你单位连续欠付工资。", "工资差额至今未补足。", "上述事实有银行流水证明。"],
        )
        indented = [
            paragraph
            for paragraph in paragraphs
            if paragraph.text in {
                "本人正式通知你单位解除劳动合同。",
                "解除自【待补：解除生效日期】生效。",
                "请依法办理结算。",
                "一、结清工资。",
                "二、出具解除证明。",
            }
        ]
        self.assertEqual(len(indented), 5)
        self.assertTrue(all(_indent_value(paragraph, "firstLineChars") == "200" for paragraph in indented))

        signature = [
            paragraph
            for paragraph in paragraphs
            if paragraph.style.name == "ArbiBuddy Legal Signature"
        ]
        self.assertEqual(len(signature), 3)
        self.assertTrue(all(_indent_value(paragraph, "leftChars") == "1600" for paragraph in signature))
        self.assertEqual(texts[texts.index(signature[0].text) - 1], "")
        self.assertEqual(sum("【待补：劳动者签名】" in text for text in texts), 1)
        self.assertFalse(any("并保留送达凭证" in text for text in texts))


if __name__ == "__main__":
    unittest.main()
